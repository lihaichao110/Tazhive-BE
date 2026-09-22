import csv
from collections.abc import Iterable
from datetime import date, datetime, time
from io import StringIO
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.services.wiki.ingestion.parsers.models import ParsedSource, SourceBatch

SUPPORTED_SUFFIXES = {".xlsx", ".xls", ".csv", ".tsv"}
TableRow = tuple[int, list[str]]
TableSection = tuple[str, list[TableRow]]


def _format_cell(value: object) -> str:
    """将单元格稳定地转换为单行文本。"""
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        text = value.isoformat()
    elif isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    else:
        text = str(value)
    return " ".join(text.replace("\t", " ").splitlines()).strip()


def _trim_row(row: Iterable[object]) -> list[str]:
    cells = [_format_cell(value) for value in row]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _load_xlsx(path: Path) -> list[TableSection]:
    workbook = None
    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        sections: list[TableSection] = []
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            rows = [
                (row_number, cells)
                for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1)
                if any(cells := _trim_row(row))
            ]
            if rows:
                sections.append((_format_cell(worksheet.title), rows))
        return sections
    except (BadZipFile, InvalidFileException, ParseError, EOFError, KeyError, OSError) as exc:
        raise ValueError("Excel 文件损坏、已加密或格式无效") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _format_xls_cell(workbook: xlrd.book.Book, cell: xlrd.sheet.Cell) -> str:
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return ""
    if cell.ctype == xlrd.XL_CELL_DATE:
        return _format_cell(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return f"#ERROR({int(cell.value)})"
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return str(int(cell.value))
    return _format_cell(cell.value)


def _load_xls(path: Path) -> list[TableSection]:
    workbook = None
    try:
        workbook = xlrd.open_workbook(path, on_demand=True)
        sections: list[TableSection] = []
        for sheet in workbook.sheets():
            if getattr(sheet, "visibility", 0) != 0:
                continue
            rows: list[TableRow] = []
            for row_index in range(sheet.nrows):
                row = [_format_xls_cell(workbook, sheet.cell(row_index, col)) for col in range(sheet.ncols)]
                while row and not row[-1]:
                    row.pop()
                if any(row):
                    rows.append((row_index + 1, row))
            if rows:
                sections.append((_format_cell(sheet.name), rows))
        return sections
    except (xlrd.XLRDError, EOFError, OSError, ValueError) as exc:
        raise ValueError("Excel 文件损坏、已加密或格式无效") from exc
    finally:
        if workbook is not None:
            workbook.release_resources()


def _decode_csv(path: Path) -> str:
    content = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 文件编码不受支持，请使用 UTF-8 或 GB18030")


def _load_csv(path: Path) -> list[TableSection]:
    text = _decode_csv(path)
    if not text.strip():
        return []
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    rows = [
        (row_number, cells)
        for row_number, row in enumerate(
            csv.reader(StringIO(text), delimiter=delimiter), start=1
        )
        if any(cells := _trim_row(row))
    ]
    return [("数据", rows)] if rows else []


def _escape_markdown_cell(value: str) -> str:
    return value.replace("\\", "\\\\").replace("|", "\\|")


def _render_rows(section_name: str, numbered_rows: list[tuple[int, list[str]]]) -> str:
    column_count = max((len(row) for _, row in numbered_rows), default=0)
    headers = ["原始行号", *(f"列{index}" for index in range(1, column_count + 1))]
    lines = [
        f"## 工作表：{section_name}",
        "",
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row_number, row in numbered_rows:
        padded = [*row, *("" for _ in range(column_count - len(row)))]
        values = [str(row_number), *(_escape_markdown_cell(value) for value in padded)]
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def _split_section(
    section_name: str,
    rows: list[TableRow],
    *,
    batch_chars: int,
) -> list[SourceBatch]:
    batches: list[SourceBatch] = []
    current: list[tuple[int, list[str]]] = []

    def append_current() -> None:
        if not current:
            return
        markdown = _render_rows(section_name, current)
        batches.append(
            SourceBatch(
                section_name=section_name,
                part_number=len(batches) + 1,
                row_start=current[0][0],
                row_end=current[-1][0],
                text=markdown,
                detail_markdown=markdown,
            )
        )

    for row_number, row in rows:
        candidate = [*current, (row_number, row)]
        if current and len(_render_rows(section_name, candidate)) > batch_chars:
            append_current()
            current = [(row_number, row)]
        else:
            current = candidate
    append_current()
    return batches


def parse_spreadsheet(path: Path, *, batch_chars: int) -> ParsedSource:
    """解析表格并按完整行切分，返回可编译批次与无损明细 Markdown。"""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        sections = _load_xlsx(path)
    elif suffix == ".xls":
        sections = _load_xls(path)
    elif suffix in {".csv", ".tsv"}:
        sections = _load_csv(path)
    else:
        raise ValueError(f"不支持的表格文件类型: {suffix}")

    batches = tuple(
        batch
        for section_name, rows in sections
        for batch in _split_section(section_name, rows, batch_chars=batch_chars)
    )
    if not batches:
        raise ValueError("表格文件中没有可提取的有效内容")
    return ParsedSource(kind="table", batches=batches)
