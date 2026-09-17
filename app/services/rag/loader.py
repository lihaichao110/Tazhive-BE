from datetime import date, datetime, time
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader, TextLoader
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class DocumentParseError(ValueError):
    """表示上传文件本身无法解析，调用方应将其视为客户端输入错误。"""


def _format_excel_cell(value: object) -> str:
    """将 Excel 单元格转换为稳定的单行文本，避免破坏表格行列结构。"""
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        text = value.isoformat()
    else:
        text = str(value)
    return (
        text.replace("\t", " ").replace("\r\n", " ").replace("\r", " ").replace("\n", " ").strip()
    )


# 表头候选单元格的最大长度：超过该长度的单元格更可能是数据（如长文本、URL），
# 用于区分真表头行和被首行标题遮挡后的普通数据行。
_HEADER_CELL_MAX_LENGTH = 30


def _header_like(row: list[str]) -> bool:
    """判断一行是否像表头：至少两个非空单元格，且都是短文本。"""
    non_empty = [cell for cell in row if cell]
    return len(non_empty) >= 2 and all(len(cell) <= _HEADER_CELL_MAX_LENGTH for cell in non_empty)


def _format_row_record(header: list[str] | None, row: list[str], sheet_name: str) -> str:
    """把一条数据行格式化为“列名: 值”的自包含记录，空单元格跳过。"""
    fields: list[str] = []
    for idx, value in enumerate(row):
        if not value:
            continue
        column = ""
        if header is not None and idx < len(header):
            column = header[idx].strip()
        if not column:
            column = f"列{idx + 1}"
        fields.append(f"{column}: {value}")
    if not fields:
        return ""
    return f"【工作表: {sheet_name}】" + "；".join(fields)


def _load_excel(path: Path) -> list[str]:
    """加载所有可见工作表，每条数据行输出一条“列名: 值”的自包含记录。

    表格数据按行建向量（一行一记录），避免整表拼成大段文本后被通用切片
    在句读处截断，导致人名等关键字段与所在行上下文分离、检索时被稀释。
    """
    workbook = None
    try:
        # data_only=False 保留公式表达式；keep_links=False 避免加载外部工作簿链接。
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        texts: list[str] = []
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue

            rows: list[list[str]] = []
            for row in worksheet.iter_rows(values_only=True):
                cells = [_format_excel_cell(value) for value in row]
                # 清除尾部空单元格，同时保留行首空单元格以维持列位置。
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append(cells)

            if not rows:
                continue

            # 首行只有一个非空单元格时视为工作表标题行，不参与数据。
            if len(rows) > 1 and len([cell for cell in rows[0] if cell]) == 1:
                rows = rows[1:]

            # 只有一行内容时无法区分表头与数据，按无表头处理。
            if len(rows) > 1 and _header_like(rows[0]):
                header: list[str] | None = rows[0]
                data_rows = rows[1:]
            else:
                header = None
                data_rows = rows

            sheet_name = _format_excel_cell(worksheet.title)
            for row in data_rows:
                record = _format_row_record(header, row, sheet_name)
                if record:
                    texts.append(record)

        if not texts:
            raise DocumentParseError("Excel 文件中没有可提取的有效内容")
        return texts
    except DocumentParseError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        ParseError,
        EOFError,
        KeyError,
        OSError,
        ValueError,
    ) as exc:
        raise DocumentParseError("Excel 文件损坏、已加密或格式无效") from exc
    finally:
        if workbook is not None:
            workbook.close()


def load_document(file_path: str) -> list[str]:
    """加载文档并返回纯文本列表（按页或整个）"""
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext in [".txt", ".md"]:
        docs = TextLoader(str(path), encoding="utf-8").load()
        return [doc.page_content for doc in docs]
    elif ext == ".pdf":
        docs = PyPDFLoader(str(path)).load()
        return [doc.page_content for doc in docs]
    elif ext == ".docx":
        docs = Docx2txtLoader(str(path)).load()
        return [doc.page_content for doc in docs]
    elif ext == ".xlsx":
        return _load_excel(path)
    else:
        raise ValueError(f"不支持的文件类型: {ext}")
