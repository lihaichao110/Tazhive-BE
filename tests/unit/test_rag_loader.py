from datetime import date, datetime

import pytest
from openpyxl import Workbook

from app.services.rag.loader import DocumentParseError, load_document


def _save_workbook(workbook, tmp_path, name="sample.xlsx"):
    file_path = tmp_path / name
    workbook.save(file_path)
    workbook.close()
    return file_path


def test_load_xlsx_formats_each_data_row_as_self_contained_record(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    sheet.append(["单位名称", "职称", "员工姓名", "分机号", "负责业务", "更新时间", "计算值"])
    sheet.append(
        [
            "应用开发一处",
            None,
            "A\n产品",
            12,
            date(2026, 9, 11),
            datetime(2026, 9, 11, 8, 30),
            "=B2*2",
        ]
    )
    sheet.append([None, None, None, None, None, None, None])

    texts = load_document(str(_save_workbook(workbook, tmp_path)))

    assert texts == [
        "【工作表: 汇总】单位名称: 应用开发一处；员工姓名: A 产品；分机号: 12；"
        "负责业务: 2026-09-11T00:00:00；更新时间: 2026-09-11T08:30:00；计算值: =B2*2"
    ]


def test_load_xlsx_drops_title_row_and_uses_next_row_as_header(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "版本记录"
    sheet.append(["版本记录"])
    sheet.append(["日期", "变更类型", "人员", "变更内容"])
    sheet.append(["44930", "creat", "郎琳", "创建文档"])

    texts = load_document(str(_save_workbook(workbook, tmp_path)))

    assert texts == [
        "【工作表: 版本记录】日期: 44930；变更类型: creat；人员: 郎琳；变更内容: 创建文档"
    ]


def test_load_xlsx_falls_back_to_generic_columns_without_header(tmp_path):
    """标题行后紧跟长文本数据行（无真表头）时，用“列N”兜底且不丢数据。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "联系方式"
    sheet.append(["各地监管局、保险行业协会联系方式"])
    sheet.append(["北京监管局网站：http://beijing.circ.gov.cn", "电话：010-66060530"])
    sheet.append(["山东监管局网站：http://shandong.circ.gov.cn"])

    texts = load_document(str(_save_workbook(workbook, tmp_path)))

    assert texts == [
        "【工作表: 联系方式】列1: 北京监管局网站：http://beijing.circ.gov.cn；列2: 电话：010-66060530",
        "【工作表: 联系方式】列1: 山东监管局网站：http://shandong.circ.gov.cn",
    ]


def test_load_xlsx_single_row_sheet_uses_generic_columns(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append([None, "保留列位置", "包含\t制表符"])

    texts = load_document(str(_save_workbook(workbook, tmp_path)))

    assert texts == ["【工作表: 明细】列2: 保留列位置；列3: 包含 制表符"]


def test_load_xlsx_skips_hidden_sheets(tmp_path):
    workbook = Workbook()
    workbook.active.title = "可见页"
    workbook.active.append(["名称", "数量"])
    workbook.active.append(["A", 1])
    hidden_sheet = workbook.create_sheet("内部数据")
    hidden_sheet.append(["不应被提取"])
    hidden_sheet.sheet_state = "hidden"

    texts = load_document(str(_save_workbook(workbook, tmp_path)))

    assert texts == ["【工作表: 可见页】名称: A；数量: 1"]


def test_load_xlsx_empty_header_cell_gets_generic_column_name(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "缺表头"
    sheet.append(["名称", "", "数量"])
    sheet.append(["A", "备注", 3])

    texts = load_document(str(_save_workbook(workbook, tmp_path)))

    assert texts == ["【工作表: 缺表头】名称: A；列2: 备注；数量: 3"]


def test_load_xlsx_rejects_workbook_without_visible_content(tmp_path):
    workbook = Workbook()
    workbook.active.title = "空白页"
    hidden_sheet = workbook.create_sheet("隐藏内容")
    hidden_sheet.append(["不参与解析"])
    hidden_sheet.sheet_state = "hidden"

    with pytest.raises(DocumentParseError, match="没有可提取的有效内容"):
        load_document(str(_save_workbook(workbook, tmp_path, "empty.xlsx")))


def test_load_xlsx_rejects_invalid_file(tmp_path):
    file_path = tmp_path / "broken.xlsx"
    file_path.write_bytes(b"not an excel workbook")

    with pytest.raises(DocumentParseError, match="损坏、已加密或格式无效"):
        load_document(str(file_path))
