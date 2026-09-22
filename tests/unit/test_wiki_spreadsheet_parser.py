from datetime import date
from types import SimpleNamespace

import pytest
import xlrd
from openpyxl import Workbook

from app.services.wiki.ingestion.parsers import parse_source, spreadsheet


def _save_workbook(workbook: Workbook, tmp_path, name: str = "sample.xlsx"):
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    return path


def test_parse_xlsx_preserves_visible_rows_and_formula(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "职业|代码"
    sheet.append(["代码", "名称", "日期", "公式"])
    sheet.append(["001", "管理\n人员", date(2026, 9, 21), "=1+1"])
    hidden = workbook.create_sheet("隐藏")
    hidden.sheet_state = "hidden"
    hidden.append(["不应出现"])

    parsed = parse_source(_save_workbook(workbook, tmp_path), batch_chars=24_000)

    assert parsed.kind == "table"
    assert len(parsed.batches) == 1
    markdown = parsed.batches[0].detail_markdown
    assert markdown is not None
    assert "## 工作表：职业|代码" in markdown
    assert "管理 人员" in markdown
    assert "2026-09-21T00:00:00" in markdown
    assert "=1+1" in markdown
    assert "隐藏" not in markdown
    assert "职业\\|代码" not in markdown  # 工作表标题不是 Markdown 表格单元格


def test_parse_xlsx_batches_only_between_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["编号", "说明"])
    for index in range(1, 5):
        sheet.append([index, f"第{index}行-" + "内容" * 10])

    parsed = parse_source(_save_workbook(workbook, tmp_path), batch_chars=150)

    assert len(parsed.batches) > 1
    combined = "\n".join(batch.detail_markdown or "" for batch in parsed.batches)
    for index in range(1, 5):
        assert combined.count(f"第{index}行-") == 1
    assert all(batch.text.startswith("## 工作表：数据") for batch in parsed.batches)
    assert all(batch.row_start is not None and batch.row_end is not None for batch in parsed.batches)


def test_parse_xlsx_keeps_physical_row_numbers_after_blank_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["代码"])
    sheet.append([])
    sheet.append(["001"])

    parsed = parse_source(_save_workbook(workbook, tmp_path), batch_chars=24_000)

    markdown = parsed.batches[0].detail_markdown or ""
    assert "| 1 | 代码 |" in markdown
    assert "| 3 | 001 |" in markdown


def test_parse_csv_supports_gb18030_delimiter_quotes_and_newlines(tmp_path):
    path = tmp_path / "codes.csv"
    path.write_bytes('代码;名称;说明\n001;"管理|人员";"第一行\n第二行"\n'.encode("gb18030"))

    parsed = parse_source(path)
    markdown = parsed.batches[0].detail_markdown or ""

    assert "管理\\|人员" in markdown
    assert "第一行 第二行" in markdown


def test_parse_tsv_uses_tab_delimiter(tmp_path):
    path = tmp_path / "codes.tsv"
    path.write_text("代码\t名称\n001\t管理人员\n", encoding="utf-8")

    parsed = parse_source(path)

    assert "| 2 | 001 | 管理人员 |" in (parsed.batches[0].detail_markdown or "")


def test_parse_xls_uses_cached_values_and_skips_hidden_sheets(tmp_path, monkeypatch):
    class FakeSheet:
        def __init__(self, name, rows, visibility=0):
            self.name = name
            self._rows = rows
            self.visibility = visibility
            self.nrows = len(rows)
            self.ncols = max(len(row) for row in rows)

        def cell(self, row, col):
            if col >= len(self._rows[row]):
                return SimpleNamespace(ctype=xlrd.XL_CELL_EMPTY, value="")
            return self._rows[row][col]

    cells = [
        [
            SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value="代码"),
            SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value="名称"),
        ],
        [
            SimpleNamespace(ctype=xlrd.XL_CELL_NUMBER, value=1.0),
            SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value="管理人员"),
        ],
    ]
    workbook = SimpleNamespace(
        datemode=0,
        sheets=lambda: [FakeSheet("职业", cells), FakeSheet("隐藏", cells, visibility=1)],
        release_resources=lambda: None,
    )
    monkeypatch.setattr(spreadsheet.xlrd, "open_workbook", lambda *_args, **_kwargs: workbook)
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"mocked")

    parsed = parse_source(path)
    markdown = parsed.batches[0].detail_markdown or ""

    assert "| 2 | 1 | 管理人员 |" in markdown
    assert "隐藏" not in markdown


@pytest.mark.parametrize("name", ["empty.xlsx", "empty.csv"])
def test_parse_spreadsheet_rejects_empty_content(tmp_path, name):
    path = tmp_path / name
    if path.suffix == ".xlsx":
        workbook = Workbook()
        _save_workbook(workbook, tmp_path, name)
    else:
        path.write_text("", encoding="utf-8")

    with pytest.raises(ValueError, match="没有可提取"):
        parse_source(path)
